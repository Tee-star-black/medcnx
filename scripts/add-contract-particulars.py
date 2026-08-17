from pathlib import Path

from docx import Document
from docx.shared import Pt


FIELDS = [
    ("Position Title", "{{document.positionTitle}}"),
    ("Commencement Date", "{{document.commencementDate}}"),
    ("Reporting Line", "{{document.reportingLine}}"),
    ("Nature of Employment", "{{document.natureOfEmployment}}"),
    ("Hours of Work", "{{document.hoursOfWork}}"),
    ("Remuneration Package", "{{document.remunerationPackage}}"),
    ("Regulatory Compliance", "{{document.regulatoryCompliance}}"),
]


def add_particulars(path: Path) -> None:
    document = Document(path)
    if any("{{document.reportingLine}}" in p.text for p in document.paragraphs):
        return

    target = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.strip().startswith("1.")
        ),
        document.paragraphs[-1],
    )

    heading = target.insert_paragraph_before()
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(4)
    run = heading.add_run("EMPLOYMENT PARTICULARS")
    run.bold = True

    for label, placeholder in FIELDS:
        paragraph = target.insert_paragraph_before()
        paragraph.paragraph_format.space_after = Pt(2)
        label_run = paragraph.add_run(f"{label}: ")
        label_run.bold = True
        paragraph.add_run(placeholder)

    spacer = target.insert_paragraph_before()
    spacer.paragraph_format.space_after = Pt(4)
    document.save(path)


root = Path("apps/api/resources/document-templates")
for contract in sorted(root.glob("*Contract*.docx")):
    add_particulars(contract)
    print(contract)
