from pathlib import Path

from docx import Document


REPLACEMENTS = {
    "[Insert Reg Number]": "{{organisation.registrationNumber}}",
    "[Insert SARS PAYE No]": "{{document.payeReferenceNumber}}",
    "[Insert Registration No / N/A]": "{{document.professionalRegistrationNumber}}",
    "[Insert Month & Year]": "{{document.payPeriod}}",
    "[Insert Bank & Acc No (Masked)]": "{{document.bankAccountMasked}}",
    "[Insert Position Title - e.g., Medical Doctor / Triage Nurse / Front Desk Admin]": "{{document.positionTitle}}",
    "[Insert Start Date, e.g., 01 July 2026]": "{{document.commencementDate}}",
    "[Insert Manager Title, e.g., Clinical Director / Practice Manager / COO]": "{{document.reportingLine}}",
    "R [Insert Monthly Salary Amount]": "{{document.remunerationPackage}}",
    "Full-Time, Permanent (subject to a standard three [3] month probationary period as governed by the Labour Relations Act)": "{{document.natureOfEmployment}}",
    "Forty (40) ordinary hours per week, scheduled dynamically in accordance with the Company's operational shift rosters (including rotational Saturdays and standby parameters)": "{{document.hoursOfWork}}",
    "Subject to providing valid proof of active, unendorsed professional registration with your respective oversight council (HPCSA / SANC) and valid Malpractice Indemnity Cover where applicable.": "{{document.regulatoryCompliance}}",
    "[Insert Details, e.g., None / Variable Shift Allowance / N/A]": "{{document.allowancesDetails}}",
}


def paragraphs(container):
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from paragraphs(cell)


def replace_in_paragraph(paragraph):
    current = "".join(run.text for run in paragraph.runs)
    updated = current
    for source, target in REPLACEMENTS.items():
        updated = updated.replace(source, target)
    if updated == current:
        return False
    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(updated)
    return True


root = Path("apps/api/resources/document-templates")
for file_path in sorted(root.glob("*.docx")):
    document = Document(file_path)
    changed = False
    for paragraph in paragraphs(document):
        changed = replace_in_paragraph(paragraph) or changed
    if changed:
        document.save(file_path)
        print(file_path)
