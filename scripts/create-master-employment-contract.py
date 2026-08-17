#!/usr/bin/env python3
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "apps/api/resources/document-templates/MedCNX_Employment_Contract_Master.docx"
NAVY, TEAL, MUTED, LIGHT = "163042", "08788D", "607787", "E8EEF2"


def font(run, size=10.5, bold=False, color=NAVY):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size, run.bold = Pt(size), bold
    run.font.color.rgb = RGBColor.from_string(color)


def border_bottom(paragraph, color=TEAL, size="10"):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", size), ("space", "6"), ("color", color)):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    ppr.append(borders)


def page_field(paragraph):
    run = paragraph.add_run()
    for kind in ("begin", "separate", "end"):
        if kind == "separate":
            instruction = OxmlElement("w:instrText")
            instruction.set(qn("xml:space"), "preserve")
            instruction.text = " PAGE "
            run._r.append(instruction)
        node = OxmlElement("w:fldChar")
        node.set(qn("w:fldCharType"), kind)
        run._r.append(node)
    font(run, 8, color=MUTED)


def heading(doc, number, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    font(p.add_run(f"{number}. {title.upper()}"), 12.5, True, TEAL)
    return p


def body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    font(p.add_run(text))
    return p


def shade(cell, fill=LIGHT):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def cell_margins(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, value in (("top", 100), ("start", 130), ("bottom", 100), ("end", 130)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcpr.append(mar)


def particulars(doc):
    rows = [
        ("Employee", "{{employee.fullName}}"),
        ("Identity number", "{{employee.idNumber}}"),
        ("Position title", "{{document.positionTitle}}"),
        ("Department", "{{employee.department}}"),
        ("Reporting line", "{{document.reportingLine}}"),
        ("Contract type", "{{document.contractType}}"),
        ("Work schedule", "{{document.employmentSchedule}}"),
        ("Commencement date", "{{document.commencementDate}}"),
        ("End date", "{{document.endDate}}"),
        ("Probation", "{{document.probationPeriod}}"),
        ("Work location", "{{document.workLocation}}"),
        ("Hours of work", "{{document.hoursOfWork}}"),
        ("Remuneration", "{{document.remunerationPackage}}"),
        ("Professional council", "{{document.professionalCouncil}}"),
        ("Registration number", "{{document.professionalRegistrationNumber}}"),
        ("Notice period", "{{document.noticePeriod}}"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    table.columns[0].width, table.columns[1].width = Inches(2.0), Inches(4.35)
    for index, (label, value) in enumerate(rows):
        left, right = table.rows[index].cells
        left.width, right.width = Inches(2.0), Inches(4.35)
        shade(left)
        for cell in (left, right):
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        font(left.paragraphs[0].add_run(label), 9.5, True)
        font(right.paragraphs[0].add_run(value), 9.5)
    return table


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(0.82), Inches(0.78)
    section.left_margin = section.right_margin = Inches(0.9)
    section.header_distance = section.footer_distance = Inches(0.42)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(6), 1.1

    h = section.header.paragraphs[0]
    h.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    font(h.add_run("{{organisation.logo}}"), 8)
    font(h.add_run("\t{{organisation.name}}"), 9, True)
    border_bottom(h)
    f = section.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(f.add_run("{{organisation.name}}  •  {{organisation.phone}}  •  {{organisation.email}}\nConfidential  |  Page "), 7.5, color=MUTED)
    page_field(f)

    title = doc.add_paragraph()
    title.paragraph_format.space_before, title.paragraph_format.space_after = Pt(18), Pt(3)
    font(title.add_run("EMPLOYMENT CONTRACT"), 21, True)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(16)
    border_bottom(sub, TEAL, "8")
    font(sub.add_run("A controlled agreement between the organisation and employee"), 10, color=MUTED)

    intro = body(doc, "This Employment Contract is concluded between {{organisation.name}} (the “Employer”) and {{employee.fullName}} (the “Employee”) on {{document.generatedDate}}.")
    intro.paragraph_format.space_after = Pt(12)
    heading(doc, 1, "Employment particulars")
    particulars(doc)

    heading(doc, 2, "Appointment and duties")
    body(doc, "The Employer appoints the Employee as {{document.positionTitle}}. The Employee will perform the duties reasonably associated with the position, follow lawful instructions, maintain professional standards and report to {{document.reportingLine}}.")
    heading(doc, 3, "Term and probation")
    body(doc, "Employment commences on {{document.commencementDate}} under {{document.contractType}} terms. Where the appointment is fixed-term, it ends on {{document.endDate}} unless lawfully renewed in writing. The probation period is {{document.probationPeriod}} and will be managed fairly in accordance with applicable labour law.")
    heading(doc, 4, "Place and hours of work")
    body(doc, "The normal place of work is {{document.workLocation}}. The agreed schedule is {{document.employmentSchedule}} and the ordinary working arrangement is {{document.hoursOfWork}}. Reasonable operational changes will be communicated in accordance with law and company policy.")
    heading(doc, 5, "Remuneration and benefits")
    body(doc, "The Employee’s remuneration package is {{document.remunerationPackage}}. Statutory deductions will be made, and any benefits or allowances are governed by the applicable payroll records, benefit rules and written company policies.")
    heading(doc, 6, "Leave")
    body(doc, "Annual, sick, family-responsibility, maternity, parental and other leave will be administered in accordance with the Basic Conditions of Employment Act, applicable collective arrangements and the Employer’s approved leave policy.")
    heading(doc, 7, "Professional and regulatory compliance")
    body(doc, "Professional council: {{document.professionalCouncil}}. Registration number: {{document.professionalRegistrationNumber}}. The Employee must maintain every registration, licence and credential required for the role. Additional requirements: {{document.regulatoryCompliance}}.")
    heading(doc, 8, "Confidentiality, information security and POPIA")
    body(doc, "The Employee must protect confidential business, employee and patient information; use systems only for authorised purposes; follow information-security controls; and comply with POPIA, healthcare confidentiality obligations and company policy during and after employment.")
    heading(doc, 9, "Conduct, performance and policies")
    body(doc, "The Employee agrees to the Employer’s lawful policies, ethical standards, performance processes, health and safety requirements, disciplinary code and grievance procedures, as amended through fair and lawful processes.")
    heading(doc, 10, "Termination")
    body(doc, "Either party may terminate employment on {{document.noticePeriod}} notice, subject to applicable law. Nothing in this agreement limits lawful termination for misconduct, incapacity, operational requirements, expiry of a fixed term or another recognised legal ground.")
    heading(doc, 11, "Special conditions")
    body(doc, "{{document.specialConditions}}")
    heading(doc, 12, "Entire agreement and governing law")
    body(doc, "This contract, together with incorporated policies and written annexures, records the employment agreement. Changes must be recorded in writing and authorised. South African law governs this agreement. If a provision is invalid, the remaining provisions continue to apply.")

    heading(doc, 13, "Signatures")
    body(doc, "The parties confirm that they have read, understood and voluntarily accepted this agreement.")
    sig = doc.add_table(rows=1, cols=2)
    sig.alignment, sig.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    for cell in sig.rows[0].cells:
        cell.width = Inches(3.15)
        cell_margins(cell)
    left, right = sig.rows[0].cells
    font(left.paragraphs[0].add_run("{{ceo.signature}}\n{{ceo.fullName}}\n{{ceo.jobTitle}}\nFor {{organisation.name}}\nDate: ____________________"), 9.5, True)
    font(right.paragraphs[0].add_run("\n____________________________\n{{employee.fullName}}\nEmployee\nDate: ____________________"), 9.5, True)

    doc.core_properties.title = "MedCNX Master Employment Contract"
    doc.core_properties.author = "MedCNX"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
